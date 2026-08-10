"""
OP System — RAG (Retrieval-Augmented Generation) для анализа документов.

Поддерживаемые форматы: PDF, DOCX, XLSX, TXT, CSV, ZIP, RAR.
Генерация ответов — через DeepSeek Chat API.
Поиск релевантных фрагментов — TF-IDF (локально, без внешних эмбеддингов).
"""

from __future__ import annotations

import csv
import os
import re
import tempfile
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import requests
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity


DEEPSEEK_API_URL = "https://api.deepseek.com/chat/completions"
DEEPSEEK_MODEL = "deepseek-chat"

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".xlsx",
    ".xls",
    ".txt",
    ".md",
    ".csv",
    ".zip",
    ".rar",
}


@dataclass
class Chunk:
    text: str
    source: str
    meta: dict[str, Any] = field(default_factory=dict)


class RAGSystem:
    """Система загрузки документов и ответов на вопросы по их содержимому."""

    def __init__(
        self,
        api_key: str | None = None,
        chunk_size: int = 800,
        chunk_overlap: int = 150,
        top_k: int = 5,
        model: str = DEEPSEEK_MODEL,
    ) -> None:
        self.api_key = api_key or os.getenv("DEEPSEEK_API_KEY", "")
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap
        self.top_k = top_k
        self.model = model

        self.chunks: list[Chunk] = []
        self._vectorizer: TfidfVectorizer | None = None
        self._matrix = None

    # ------------------------------------------------------------------
    # Публичный API
    # ------------------------------------------------------------------

    def load_file(self, file_path: str) -> dict[str, Any]:
        """Загружает файл, разбивает на чанки и индексирует."""
        path = Path(file_path).expanduser().resolve()
        if not path.exists():
            raise FileNotFoundError(f"Файл не найден: {path}")

        ext = path.suffix.lower()
        if ext not in SUPPORTED_EXTENSIONS:
            raise ValueError(
                f"Неподдерживаемый формат: {ext}. "
                f"Допустимо: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
            )

        texts = self._extract_texts(path)
        if not texts:
            return {"chunks_count": 0, "sources": []}

        new_chunks: list[Chunk] = []
        for source_name, text in texts:
            for i, piece in enumerate(self._chunk_text(text)):
                new_chunks.append(
                    Chunk(
                        text=piece,
                        source=source_name,
                        meta={"chunk_index": i, "parent": path.name},
                    )
                )

        self.chunks.extend(new_chunks)
        self._rebuild_index()

        return {
            "chunks_count": len(new_chunks),
            "total_chunks": len(self.chunks),
            "sources": sorted({c.source for c in new_chunks}),
        }

    def ask(self, question: str) -> dict[str, Any]:
        """Ищет релевантные фрагменты и генерирует ответ через DeepSeek."""
        question = (question or "").strip()
        if not question:
            return {"answer": "Пустой вопрос.", "sources": []}

        if not self.chunks:
            return {
                "answer": "Документы ещё не загружены. Сначала вызовите load_file().",
                "sources": [],
            }

        sources = self._retrieve(question, top_k=self.top_k)
        context = "\n\n---\n\n".join(
            f"[Источник: {s['source']}]\n{s['text']}" for s in sources
        )

        answer = self._generate_answer(question, context)
        return {"answer": answer, "sources": sources}

    def clear(self) -> None:
        """Очищает загруженные документы и индекс."""
        self.chunks = []
        self._vectorizer = None
        self._matrix = None

    # ------------------------------------------------------------------
    # Извлечение текста
    # ------------------------------------------------------------------

    def _extract_texts(self, path: Path) -> list[tuple[str, str]]:
        ext = path.suffix.lower()
        if ext == ".pdf":
            return [(path.name, self._read_pdf(path))]
        if ext == ".docx":
            return [(path.name, self._read_docx(path))]
        if ext in {".xlsx", ".xls"}:
            return [(path.name, self._read_xlsx(path))]
        if ext in {".txt", ".md"}:
            return [(path.name, path.read_text(encoding="utf-8", errors="ignore"))]
        if ext == ".csv":
            return [(path.name, self._read_csv(path))]
        if ext == ".zip":
            return self._read_archive_zip(path)
        if ext == ".rar":
            return self._read_archive_rar(path)
        return []

    @staticmethod
    def _read_pdf(path: Path) -> str:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        parts: list[str] = []
        for i, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            text = text.strip()
            if text:
                parts.append(f"[Страница {i}]\n{text}")
        return "\n\n".join(parts)

    @staticmethod
    def _read_docx(path: Path) -> str:
        from docx import Document

        doc = Document(str(path))
        parts = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)

    @staticmethod
    def _read_xlsx(path: Path) -> str:
        from openpyxl import load_workbook

        wb = load_workbook(str(path), data_only=True, read_only=True)
        parts: list[str] = []
        for sheet in wb.worksheets:
            parts.append(f"[Лист: {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                if cells:
                    parts.append(" | ".join(cells))
        wb.close()
        return "\n".join(parts)

    @staticmethod
    def _read_csv(path: Path) -> str:
        parts: list[str] = []
        with path.open("r", encoding="utf-8", errors="ignore", newline="") as f:
            reader = csv.reader(f)
            for row in reader:
                cells = [c.strip() for c in row if c and c.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)

    def _read_archive_zip(self, path: Path) -> list[tuple[str, str]]:
        results: list[tuple[str, str]] = []
        with zipfile.ZipFile(path, "r") as zf:
            for info in zf.infolist():
                if info.is_dir():
                    continue
                inner = Path(info.filename)
                if inner.suffix.lower() not in SUPPORTED_EXTENSIONS - {".zip", ".rar"}:
                    continue
                with zf.open(info) as src, tempfile.NamedTemporaryFile(
                    suffix=inner.suffix, delete=False
                ) as tmp:
                    tmp.write(src.read())
                    tmp_path = Path(tmp.name)
                try:
                    for source_name, text in self._extract_texts(tmp_path):
                        label = f"{path.name}/{info.filename}"
                        if text.strip():
                            results.append((label, text))
                finally:
                    tmp_path.unlink(missing_ok=True)
        return results

    def _read_archive_rar(self, path: Path) -> list[tuple[str, str]]:
        try:
            import rarfile
        except ImportError as exc:
            raise RuntimeError(
                "Для RAR нужен пакет rarfile и утилита unrar в системе."
            ) from exc

        results: list[tuple[str, str]] = []
        with rarfile.RarFile(path) as rf:
            for info in rf.infolist():
                if info.is_dir():
                    continue
                inner = Path(info.filename)
                if inner.suffix.lower() not in SUPPORTED_EXTENSIONS - {".zip", ".rar"}:
                    continue
                with rf.open(info) as src, tempfile.NamedTemporaryFile(
                    suffix=inner.suffix, delete=False
                ) as tmp:
                    tmp.write(src.read())
                    tmp_path = Path(tmp.name)
                try:
                    for source_name, text in self._extract_texts(tmp_path):
                        label = f"{path.name}/{info.filename}"
                        if text.strip():
                            results.append((label, text))
                finally:
                    tmp_path.unlink(missing_ok=True)
        return results

    # ------------------------------------------------------------------
    # Чанкинг и индекс
    # ------------------------------------------------------------------

    def _chunk_text(self, text: str) -> list[str]:
        text = re.sub(r"\r\n?", "\n", text or "")
        text = re.sub(r"[ \t]+", " ", text)
        text = re.sub(r"\n{3,}", "\n\n", text).strip()
        if not text:
            return []

        if len(text) <= self.chunk_size:
            return [text]

        chunks: list[str] = []
        start = 0
        while start < len(text):
            end = min(start + self.chunk_size, len(text))
            if end < len(text):
                # стараемся резать по границе предложения/абзаца
                window = text[start:end]
                split_at = max(window.rfind("\n\n"), window.rfind(". "), window.rfind("\n"))
                if split_at > self.chunk_size // 3:
                    end = start + split_at + 1
            piece = text[start:end].strip()
            if piece:
                chunks.append(piece)
            if end >= len(text):
                break
            start = max(0, end - self.chunk_overlap)
        return chunks

    def _rebuild_index(self) -> None:
        if not self.chunks:
            self._vectorizer = None
            self._matrix = None
            return
        corpus = [c.text for c in self.chunks]
        self._vectorizer = TfidfVectorizer(
            max_features=20000,
            ngram_range=(1, 2),
            sublinear_tf=True,
        )
        self._matrix = self._vectorizer.fit_transform(corpus)

    def _retrieve(self, question: str, top_k: int = 5) -> list[dict[str, Any]]:
        if self._vectorizer is None or self._matrix is None or not self.chunks:
            return []

        q_vec = self._vectorizer.transform([question])
        scores = cosine_similarity(q_vec, self._matrix).ravel()
        ranked = scores.argsort()[::-1]

        results: list[dict[str, Any]] = []
        seen: set[str] = set()
        for idx in ranked:
            if scores[idx] <= 0:
                continue
            chunk = self.chunks[int(idx)]
            key = chunk.text[:120]
            if key in seen:
                continue
            seen.add(key)
            results.append(
                {
                    "text": chunk.text,
                    "source": chunk.source,
                    "score": float(scores[idx]),
                    "meta": chunk.meta,
                }
            )
            if len(results) >= top_k:
                break
        return results

    # ------------------------------------------------------------------
    # DeepSeek
    # ------------------------------------------------------------------

    def _generate_answer(self, question: str, context: str) -> str:
        if not self.api_key:
            return (
                "API-ключ DeepSeek не задан. Установите переменную окружения "
                "DEEPSEEK_API_KEY или передайте api_key в RAGSystem(...)."
            )

        if not context.strip():
            return "Не удалось найти релевантные фрагменты в загруженных документах."

        system_prompt = (
            "Ты аналитик документов OP System. Отвечай только на основе "
            "предоставленного контекста. Если ответа нет в контексте — прямо скажи об этом. "
            "Отвечай на языке вопроса. Будь конкретным и структурированным."
        )
        user_prompt = (
            f"Контекст из документов:\n\n{context}\n\n"
            f"Вопрос: {question}\n\n"
            "Ответ:"
        )

        try:
            response = requests.post(
                DEEPSEEK_API_URL,
                headers={
                    "Authorization": f"Bearer {self.api_key}",
                    "Content-Type": "application/json",
                },
                json={
                    "model": self.model,
                    "messages": [
                        {"role": "system", "content": system_prompt},
                        {"role": "user", "content": user_prompt},
                    ],
                    "temperature": 0.2,
                },
                timeout=90,
            )
            response.raise_for_status()
            data = response.json()
            return data["choices"][0]["message"]["content"].strip()
        except requests.HTTPError as exc:
            detail = ""
            try:
                detail = response.text[:500]
            except Exception:
                detail = str(exc)
            return f"Ошибка DeepSeek API: {exc}. {detail}"
        except Exception as exc:  # noqa: BLE001
            return f"Не удалось получить ответ от DeepSeek: {exc}"
