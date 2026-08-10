"""
🚀 OP SYSTEM — ПОЛНАЯ ВЕРСИЯ (ОДИН ФАЙЛ)
Готово к копированию и запуску.
Автоматически устанавливает зависимости, загружает документы, отвечает на вопросы.

Запуск:
    python op_system.py

API-ключ:
    export DEEPSEEK_API_KEY='sk-...'
    # или введите в консоли при старте
"""

from __future__ import annotations

import os
import re
import subprocess
import sys
import tempfile
import warnings
import zipfile
from pathlib import Path
from typing import Any

warnings.filterwarnings("ignore")

# ============================================================
# 1. АВТОУСТАНОВКА ЗАВИСИМОСТЕЙ
# ============================================================

# pip-имя → имя модуля для import
PACKAGE_IMPORTS: dict[str, str] = {
    "requests": "requests",
    "pypdf": "pypdf",
    "python-docx": "docx",
    "openpyxl": "openpyxl",
    "scikit-learn": "sklearn",
    "rarfile": "rarfile",
    "numpy": "numpy",
}

REQUIRED_PACKAGES = [
    "requests>=2.31.0",
    "pypdf>=4.0.0",
    "python-docx>=1.1.0",
    "openpyxl>=3.1.0",
    "scikit-learn>=1.3.0",
    "rarfile>=4.1",
    "numpy>=1.26.0",
]


def _package_root(spec: str) -> str:
    """'python-docx>=1.1.0' -> 'python-docx'."""
    return re.split(r"[><=!]", spec, maxsplit=1)[0].strip()


def install_package(spec: str) -> None:
    """Устанавливает пакет через pip, если модуль ещё не импортируется."""
    root = _package_root(spec)
    module_name = PACKAGE_IMPORTS.get(root, root.replace("-", "_"))
    try:
        __import__(module_name)
        return
    except ImportError:
        pass

    print(f"📦 Устанавливаю: {spec}")
    subprocess.check_call(
        [sys.executable, "-m", "pip", "install", spec, "-q"],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.STDOUT,
    )


for _pkg in REQUIRED_PACKAGES:
    install_package(_pkg)

# ============================================================
# 2. ИМПОРТЫ (после установки)
# ============================================================

from dataclasses import dataclass, field  # noqa: E402

import requests  # noqa: E402
from docx import Document  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from pypdf import PdfReader  # noqa: E402
from sklearn.feature_extraction.text import TfidfVectorizer  # noqa: E402
from sklearn.metrics.pairwise import cosine_similarity  # noqa: E402

# ============================================================
# 3. RAGSystem
# ============================================================

DEEPSEEK_API_URL = "https://api.deepseek.com/chat/completions"
DEEPSEEK_MODEL = "deepseek-chat"

SUPPORTED_INNER = {".pdf", ".docx", ".xlsx", ".xls", ".txt", ".md", ".csv"}


@dataclass
class Chunk:
    text: str
    source: str
    meta: dict = field(default_factory=dict)


class RAGSystem:
    def __init__(
        self,
        api_key: str | None = None,
        chunk_size: int = 800,
        chunk_overlap: int = 150,
        top_k: int = 5,
    ) -> None:
        self.api_key = api_key or os.getenv("DEEPSEEK_API_KEY")
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap
        self.top_k = top_k
        self.chunks: list[Chunk] = []
        self._vectorizer = None
        self._matrix = None

    def load_file(self, file_path: str) -> dict[str, Any]:
        """Загружает и индексирует файл (PDF, DOCX, XLSX, TXT, ZIP, RAR)."""
        path = Path(file_path).expanduser()
        if not path.exists():
            return {"error": f"Файл не найден: {path}"}

        print(f"📄 Обработка: {path.name}")
        texts = self._extract_texts(path)
        if not texts:
            return {"error": "Не удалось извлечь текст"}

        new_chunks: list[Chunk] = []
        for source, text in texts:
            if not (text or "").strip():
                continue
            for piece in self._chunk_text(text):
                new_chunks.append(Chunk(text=piece, source=source))

        if not new_chunks:
            return {"error": "Не удалось извлечь текст"}

        self.chunks.extend(new_chunks)
        self._rebuild_index()

        return {
            "chunks_count": len(new_chunks),
            "total": len(self.chunks),
            "sources": sorted({c.source for c in new_chunks}),
        }

    def ask(self, question: str) -> dict[str, Any]:
        """Ищет релевантные фрагменты и генерирует ответ через DeepSeek."""
        if not self.chunks:
            return {"answer": "❌ Сначала загрузите документ.", "sources": []}

        sources = self._retrieve(question, self.top_k)
        if not sources:
            return {"answer": "❌ Ничего не найдено.", "sources": []}

        context = "\n\n---\n\n".join(
            f"[Источник: {s['source']}]\n{s['text']}" for s in sources
        )
        answer = self._generate_answer(question, context)
        return {"answer": answer, "sources": sources}

    def clear(self) -> None:
        self.chunks = []
        self._vectorizer = None
        self._matrix = None

    # ---- извлечение текста ----

    def _extract_texts(self, path: Path) -> list[tuple[str, str]]:
        ext = path.suffix.lower()
        if ext == ".pdf":
            return [(path.name, self._read_pdf(path))]
        if ext == ".docx":
            return [(path.name, self._read_docx(path))]
        if ext in {".xlsx", ".xls"}:
            return [(path.name, self._read_xlsx(path))]
        if ext in {".txt", ".md", ".csv"}:
            return [(path.name, path.read_text(encoding="utf-8", errors="ignore"))]
        if ext == ".zip":
            return self._read_archive_zip(path)
        if ext == ".rar":
            return self._read_archive_rar(path)
        return []

    def _read_pdf(self, path: Path) -> str:
        try:
            reader = PdfReader(str(path))
            return "\n\n".join(
                f"[Страница {i}]\n{p.extract_text() or ''}"
                for i, p in enumerate(reader.pages, 1)
            )
        except Exception as e:  # noqa: BLE001
            print(f"   ⚠️ Ошибка PDF: {e}")
            return ""

    def _read_docx(self, path: Path) -> str:
        try:
            doc = Document(str(path))
            parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
            return "\n".join(parts)
        except Exception as e:  # noqa: BLE001
            print(f"   ⚠️ Ошибка DOCX: {e}")
            return ""

    def _read_xlsx(self, path: Path) -> str:
        try:
            wb = load_workbook(str(path), data_only=True, read_only=True)
            parts: list[str] = []
            for sheet in wb.worksheets:
                parts.append(f"[Лист: {sheet.title}]")
                for row in sheet.iter_rows(values_only=True):
                    cells = [
                        str(c).strip()
                        for c in row
                        if c is not None and str(c).strip()
                    ]
                    if cells:
                        parts.append(" | ".join(cells))
            wb.close()
            return "\n".join(parts)
        except Exception as e:  # noqa: BLE001
            print(f"   ⚠️ Ошибка Excel: {e}")
            return ""

    def _read_archive_zip(self, path: Path) -> list[tuple[str, str]]:
        results: list[tuple[str, str]] = []
        try:
            with zipfile.ZipFile(path, "r") as zf:
                for info in zf.infolist():
                    if info.is_dir():
                        continue
                    inner = Path(info.filename)
                    if inner.suffix.lower() not in SUPPORTED_INNER:
                        continue
                    with zf.open(info) as src, tempfile.NamedTemporaryFile(
                        suffix=inner.suffix, delete=False
                    ) as tmp:
                        tmp.write(src.read())
                        tmp_path = Path(tmp.name)
                    try:
                        for _source, text in self._extract_texts(tmp_path):
                            if text.strip():
                                results.append(
                                    (f"{path.name}/{info.filename}", text)
                                )
                    finally:
                        tmp_path.unlink(missing_ok=True)
            return results
        except Exception as e:  # noqa: BLE001
            print(f"   ⚠️ Ошибка ZIP: {e}")
            return []

    def _read_archive_rar(self, path: Path) -> list[tuple[str, str]]:
        try:
            import rarfile
        except ImportError:
            print("   ⚠️ Для RAR нужен пакет rarfile: pip install rarfile")
            return []

        results: list[tuple[str, str]] = []
        try:
            with rarfile.RarFile(path) as rf:
                for info in rf.infolist():
                    if info.is_dir():
                        continue
                    inner = Path(info.filename)
                    if inner.suffix.lower() not in SUPPORTED_INNER:
                        continue
                    with rf.open(info) as src, tempfile.NamedTemporaryFile(
                        suffix=inner.suffix, delete=False
                    ) as tmp:
                        tmp.write(src.read())
                        tmp_path = Path(tmp.name)
                    try:
                        for _source, text in self._extract_texts(tmp_path):
                            if text.strip():
                                results.append(
                                    (f"{path.name}/{info.filename}", text)
                                )
                    finally:
                        tmp_path.unlink(missing_ok=True)
            return results
        except Exception as e:  # noqa: BLE001
            print(f"   ⚠️ Ошибка RAR: {e}")
            return []

    # ---- чанкинг и индекс ----

    def _chunk_text(self, text: str) -> list[str]:
        text = re.sub(r"\r\n?", "\n", text or "")
        text = re.sub(r"[ \t]+", " ", text)
        text = re.sub(r"\n{3,}", "\n\n", text).strip()
        if not text:
            return []
        if len(text) <= self.chunk_size:
            return [text]

        chunks: list[str] = []
        start = 0
        while start < len(text):
            end = min(start + self.chunk_size, len(text))
            if end < len(text):
                window = text[start:end]
                split_at = max(
                    window.rfind("\n\n"),
                    window.rfind(". "),
                    window.rfind("\n"),
                )
                if split_at > self.chunk_size // 3:
                    end = start + split_at + 1
            piece = text[start:end].strip()
            if piece:
                chunks.append(piece)
            if end >= len(text):
                break
            start = max(0, end - self.chunk_overlap)
        return chunks

    def _rebuild_index(self) -> None:
        if not self.chunks:
            self._vectorizer = None
            self._matrix = None
            return
        corpus = [c.text for c in self.chunks]
        self._vectorizer = TfidfVectorizer(
            max_features=20000,
            ngram_range=(1, 2),
            sublinear_tf=True,
        )
        self._matrix = self._vectorizer.fit_transform(corpus)

    def _retrieve(self, question: str, top_k: int = 5) -> list[dict]:
        if self._vectorizer is None or self._matrix is None:
            return []

        q_vec = self._vectorizer.transform([question])
        scores = cosine_similarity(q_vec, self._matrix).ravel()
        ranked = scores.argsort()[::-1]

        results: list[dict] = []
        seen: set[str] = set()
        for idx in ranked:
            if scores[idx] <= 0:
                continue
            chunk = self.chunks[int(idx)]
            key = chunk.text[:120]
            if key in seen:
                continue
            seen.add(key)
            results.append(
                {
                    "text": chunk.text,
                    "source": chunk.source,
                    "score": float(scores[idx]),
                }
            )
            if len(results) >= top_k:
                break
        return results

    def _generate_answer(self, question: str, context: str) -> str:
        if not self.api_key:
            return (
                "❌ Ошибка: DEEPSEEK_API_KEY не задан. "
                "Установите переменную окружения."
            )
        if not context.strip():
            return "❌ Нет релевантного контента."

        system_prompt = (
            "Ты аналитик документов OP System. Отвечай строго по контексту. "
            "Если данных нет — прямо скажи. Отвечай на языке вопроса."
        )
        user_prompt = f"Контекст:\n{context}\n\nВопрос: {question}\n\nОтвет:"

        try:
            response = requests.post(
                DEEPSEEK_API_URL,
                headers={
                    "Authorization": f"Bearer {self.api_key}",
                    "Content-Type": "application/json",
                },
                json={
                    "model": DEEPSEEK_MODEL,
                    "messages": [
                        {"role": "system", "content": system_prompt},
                        {"role": "user", "content": user_prompt},
                    ],
                    "temperature": 0.2,
                },
                timeout=90,
            )
            response.raise_for_status()
            return response.json()["choices"][0]["message"]["content"].strip()
        except Exception as e:  # noqa: BLE001
            return f"⚠️ Ошибка DeepSeek: {e}"


# ============================================================
# 4. CLI
# ============================================================


def main() -> None:
    print("\n" + "=" * 60)
    print("🚀 OP SYSTEM — АНАЛИЗ ДОКУМЕНТОВ")
    print("=" * 60)

    api_key = os.getenv("DEEPSEEK_API_KEY")
    if not api_key:
        print("\n⚠️ DEEPSEEK_API_KEY не найден в переменных окружения.")
        api_key = input("🔑 Введите ваш DeepSeek API-ключ: ").strip()
        if not api_key:
            print("❌ API-ключ обязателен.")
            return
        os.environ["DEEPSEEK_API_KEY"] = api_key

    rag = RAGSystem(api_key=api_key)

    while True:
        print("\n" + "-" * 60)
        file_path = (
            input(
                "📁 Введите путь к файлу (PDF, DOCX, XLSX, ZIP, RAR) "
                "или 'exit' для выхода: "
            )
            .strip()
            .strip('"')
            .strip("'")
        )
        if file_path.lower() in {"exit", "quit", "выход"}:
            print("👋 До свидания!")
            break

        if not os.path.exists(file_path):
            print(f"❌ Файл не найден: {file_path}")
            continue

        result = rag.load_file(file_path)
        if "error" in result:
            print(f"❌ {result['error']}")
            continue

        print(f"✅ Загружено чанков: {result['chunks_count']} (Всего: {result['total']})")
        print(f"📂 Источники: {', '.join(result['sources'])}")

        while True:
            question = input(
                "\n❓ Введите вопрос (или 'назад' для выбора другого файла): "
            ).strip()
            if question.lower() in {"назад", "back", "exit", "выход"}:
                break
            if not question:
                continue

            answer = rag.ask(question)
            print("\n" + "=" * 60)
            print("🤖 ОТВЕТ:")
            print("=" * 60)
            print(answer["answer"])
            print("=" * 60)

            if answer.get("sources"):
                print("\n📚 ИСТОЧНИКИ:")
                for i, src in enumerate(answer["sources"][:3], 1):
                    preview = src["text"][:200].replace("\n", " ")
                    print(f"{i}. [{src['source']}] {preview}...")


if __name__ == "__main__":
    main()
