"""Smoke-тесты загрузки и поиска без вызова DeepSeek API."""

from __future__ import annotations

import zipfile
from pathlib import Path

from openpyxl import Workbook
from docx import Document

from rag import RAGSystem


def _make_sample_dir(tmp: Path) -> Path:
    tmp.mkdir(parents=True, exist_ok=True)

    txt = tmp / "contract.txt"
    txt.write_text(
        "Договор поставки №42.\n"
        "Стороны: ООО Альфа и ООО Бета.\n"
        "Срок поставки: 15 марта 2026 года.\n"
        "Сумма договора: 1 250 000 рублей.\n",
        encoding="utf-8",
    )

    docx_path = tmp / "report.docx"
    doc = Document()
    doc.add_heading("Отчёт по проекту OP System", level=1)
    doc.add_paragraph(
        "OP System анализирует PDF и DOCX документы и отвечает на вопросы."
    )
    doc.save(docx_path)

    xlsx_path = tmp / "budget.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Бюджет"
    ws.append(["Статья", "Сумма"])
    ws.append(["Разработка", 500000])
    ws.append(["Тестирование", 120000])
    wb.save(xlsx_path)

    zip_path = tmp / "bundle.zip"
    with zipfile.ZipFile(zip_path, "w") as zf:
        zf.write(txt, arcname="contract.txt")

    return tmp


def test_load_txt_and_ask_retrieval(tmp_path: Path) -> None:
    samples = _make_sample_dir(tmp_path)
    rag = RAGSystem(api_key="")  # без сети — проверяем только retrieval

    result = rag.load_file(str(samples / "contract.txt"))
    assert result["chunks_count"] >= 1

    sources = rag._retrieve("Какая сумма договора?", top_k=3)
    assert sources
    joined = " ".join(s["text"] for s in sources)
    assert "1 250 000" in joined or "договора" in joined.lower()


def test_load_docx_xlsx_zip(tmp_path: Path) -> None:
    samples = _make_sample_dir(tmp_path)
    rag = RAGSystem(api_key="")

    r1 = rag.load_file(str(samples / "report.docx"))
    r2 = rag.load_file(str(samples / "budget.xlsx"))
    r3 = rag.load_file(str(samples / "bundle.zip"))

    assert r1["chunks_count"] >= 1
    assert r2["chunks_count"] >= 1
    assert r3["chunks_count"] >= 1
    assert rag.chunks


if __name__ == "__main__":
    import tempfile

    root = Path(tempfile.mkdtemp(prefix="op-rag-"))
    test_load_txt_and_ask_retrieval(root / "a")
    test_load_docx_xlsx_zip(root / "b")
    print("✅ smoke tests passed")
