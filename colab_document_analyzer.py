# =============================================================================
# Система анализа документов для Google Colab
# Скопируйте ВЕСЬ этот файл в одну ячейку Colab и запустите.
# =============================================================================

# --- Установка зависимостей ---
import os
import subprocess
import sys


def _pip_install(*pkgs):
    subprocess.check_call(
        [sys.executable, "-m", "pip", "install", "-q", *pkgs],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )


_pip_install(
    "pymupdf",
    "python-docx",
    "openpyxl",
    "pytesseract",
    "pdf2image",
    "Pillow",
    "scikit-learn",
    "requests",
    "rarfile",
    "py7zr",
)

# Системные пакеты для OCR и архивов
_apt = ["apt-get"]
# В Colab обычно root; если нет — пробуем sudo
if hasattr(os, "geteuid") and os.geteuid() != 0:
    _apt = ["sudo", "apt-get"]

subprocess.run(
    _apt + ["update", "-qq"],
    check=False,
    stdout=subprocess.DEVNULL,
    stderr=subprocess.DEVNULL,
)
subprocess.run(
    _apt
    + [
        "install",
        "-y",
        "-qq",
        "tesseract-ocr",
        "tesseract-ocr-rus",
        "tesseract-ocr-eng",
        "poppler-utils",
        "unrar",
        "p7zip-full",
    ],
    check=False,
    stdout=subprocess.DEVNULL,
    stderr=subprocess.DEVNULL,
)

# --- Импорты ---
import io
import re
import zipfile
from pathlib import Path
from typing import List, Dict, Tuple, Optional

import fitz  # PyMuPDF
import pytesseract
import rarfile
import py7zr
import requests
from docx import Document
from openpyxl import load_workbook
from pdf2image import convert_from_bytes
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
from google.colab import files

# =============================================================================
# Конфигурация
# =============================================================================
API_KEY = "sk-e4d831ec99e04293829276560157b2cc"
DEEPSEEK_URL = "https://api.deepseek.com/v1/chat/completions"
DEEPSEEK_MODEL = "deepseek-chat"

CHUNK_SIZE = 800
CHUNK_OVERLAP = 150
TOP_K = 5
MIN_TEXT_CHARS_FOR_SCAN = 80  # если из PDF извлечено мало текста — запускаем OCR
WORK_DIR = Path("/content/uploaded_docs")
WORK_DIR.mkdir(parents=True, exist_ok=True)

# =============================================================================
# Извлечение текста
# =============================================================================


def extract_txt(data: bytes) -> str:
    for enc in ("utf-8", "cp1251", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def extract_docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_xlsx(data: bytes) -> str:
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    parts = []
    for sheet in wb.worksheets:
        parts.append(f"[Лист: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                parts.append(" | ".join(cells))
    wb.close()
    return "\n".join(parts)


def extract_pdf_text(data: bytes) -> str:
    parts = []
    with fitz.open(stream=data, filetype="pdf") as doc:
        for page in doc:
            text = page.get_text("text") or ""
            if text.strip():
                parts.append(text)
    return "\n".join(parts)


def extract_pdf_ocr(data: bytes, dpi: int = 200) -> str:
    """OCR для сканов через pdf2image + pytesseract."""
    images = convert_from_bytes(data, dpi=dpi)
    parts = []
    for i, img in enumerate(images, 1):
        text = pytesseract.image_to_string(img, lang="rus+eng")
        if text.strip():
            parts.append(f"[Страница {i}]\n{text}")
    return "\n".join(parts)


def extract_pdf(data: bytes) -> str:
    text = extract_pdf_text(data)
    # Если текста мало — вероятно скан, запускаем OCR
    if len(re.sub(r"\s+", "", text)) < MIN_TEXT_CHARS_FOR_SCAN:
        print("  → мало текста, запускаю OCR (pytesseract + pdf2image)...")
        ocr_text = extract_pdf_ocr(data)
        if ocr_text.strip():
            return ocr_text
    return text


def extract_image_ocr(data: bytes) -> str:
    from PIL import Image

    img = Image.open(io.BytesIO(data))
    return pytesseract.image_to_string(img, lang="rus+eng")


# =============================================================================
# Архивы
# =============================================================================

ARCHIVE_EXTS = {".zip", ".rar", ".7z"}
DOC_EXTS = {".pdf", ".docx", ".xlsx", ".txt", ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}


def _safe_name(name: str) -> str:
    return re.sub(r"[^\w.\-]+", "_", name, flags=re.UNICODE)


def unpack_zip(data: bytes, dest: Path) -> List[Path]:
    out = []
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        for info in zf.infolist():
            if info.is_dir():
                continue
            name = Path(info.filename).name
            if not name or name.startswith("."):
                continue
            target = dest / _safe_name(name)
            target.write_bytes(zf.read(info))
            out.append(target)
    return out


def unpack_rar(data: bytes, dest: Path) -> List[Path]:
    out = []
    with rarfile.RarFile(io.BytesIO(data)) as rf:
        for info in rf.infolist():
            if info.is_dir():
                continue
            name = Path(info.filename).name
            if not name or name.startswith("."):
                continue
            target = dest / _safe_name(name)
            target.write_bytes(rf.read(info))
            out.append(target)
    return out


def unpack_7z(data: bytes, dest: Path) -> List[Path]:
    out = []
    with py7zr.SevenZipFile(io.BytesIO(data), mode="r") as archive:
        archive.extractall(path=dest)
    for p in dest.rglob("*"):
        if p.is_file() and not p.name.startswith("."):
            out.append(p)
    return out


def unpack_archive(path: Path, data: bytes) -> List[Path]:
    dest = WORK_DIR / f"_extracted_{path.stem}"
    dest.mkdir(parents=True, exist_ok=True)
    ext = path.suffix.lower()
    if ext == ".zip":
        return unpack_zip(data, dest)
    if ext == ".rar":
        return unpack_rar(data, dest)
    if ext == ".7z":
        return unpack_7z(data, dest)
    return []


# =============================================================================
# Единый пайплайн извлечения
# =============================================================================


def extract_from_file(path: Path, data: Optional[bytes] = None) -> List[Tuple[str, str]]:
    """Возвращает список (source_name, text). Архивы разворачиваются рекурсивно."""
    if data is None:
        data = path.read_bytes()

    ext = path.suffix.lower()
    results: List[Tuple[str, str]] = []

    if ext in ARCHIVE_EXTS:
        print(f"  Распаковка архива: {path.name}")
        try:
            members = unpack_archive(path, data)
        except Exception as e:
            print(f"  ✗ Ошибка распаковки {path.name}: {e}")
            return results
        for member in members:
            results.extend(extract_from_file(member))
        return results

    try:
        if ext == ".pdf":
            text = extract_pdf(data)
        elif ext == ".docx":
            text = extract_docx(data)
        elif ext == ".xlsx":
            text = extract_xlsx(data)
        elif ext == ".txt":
            text = extract_txt(data)
        elif ext in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}:
            print(f"  → OCR изображения: {path.name}")
            text = extract_image_ocr(data)
        else:
            print(f"  ⚠ Пропуск неподдерживаемого файла: {path.name}")
            return results

        text = (text or "").strip()
        if text:
            results.append((path.name, text))
            print(f"  ✓ {path.name}: {len(text)} символов")
        else:
            print(f"  ⚠ {path.name}: текст не извлечён")
    except Exception as e:
        print(f"  ✗ Ошибка {path.name}: {e}")

    return results


# =============================================================================
# Чанки
# =============================================================================


def chunk_text(text: str, chunk_size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> List[str]:
    text = re.sub(r"\s+", " ", text).strip()
    if not text:
        return []
    if len(text) <= chunk_size:
        return [text]

    chunks = []
    start = 0
    step = max(1, chunk_size - overlap)
    while start < len(text):
        end = start + chunk_size
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= len(text):
            break
        start += step
    return chunks


# =============================================================================
# TF-IDF индекс
# =============================================================================


class TfidfIndex:
    def __init__(self):
        self.chunks: List[Dict] = []
        self.vectorizer: Optional[TfidfVectorizer] = None
        self.matrix = None

    def build(self, documents: List[Tuple[str, str]]):
        self.chunks = []
        for source, text in documents:
            for i, chunk in enumerate(chunk_text(text)):
                self.chunks.append(
                    {
                        "source": source,
                        "chunk_id": i,
                        "text": chunk,
                    }
                )

        if not self.chunks:
            self.vectorizer = None
            self.matrix = None
            return

        corpus = [c["text"] for c in self.chunks]
        self.vectorizer = TfidfVectorizer(
            analyzer="word",
            ngram_range=(1, 2),
            max_df=0.95,
            min_df=1,
            sublinear_tf=True,
        )
        self.matrix = self.vectorizer.fit_transform(corpus)
        print(f"Индекс: {len(self.chunks)} чанков из {len(documents)} документ(ов)")

    def search(self, query: str, top_k: int = TOP_K) -> List[Dict]:
        if not self.chunks or self.vectorizer is None or self.matrix is None:
            return []
        q = self.vectorizer.transform([query])
        scores = cosine_similarity(q, self.matrix).ravel()
        idxs = scores.argsort()[::-1][:top_k]
        results = []
        for i in idxs:
            if scores[i] <= 0:
                continue
            item = dict(self.chunks[i])
            item["score"] = float(scores[i])
            results.append(item)
        return results


# =============================================================================
# DeepSeek
# =============================================================================


def ask_deepseek(question: str, contexts: List[Dict]) -> str:
    if not contexts:
        return "Не удалось найти релевантные фрагменты в загруженных документах."

    context_blocks = []
    for i, c in enumerate(contexts, 1):
        context_blocks.append(
            f"[{i}] Источник: {c['source']} (чанк #{c['chunk_id']}, score={c['score']:.3f})\n{c['text']}"
        )
    context_text = "\n\n".join(context_blocks)

    system_prompt = (
        "Ты помощник по анализу документов. Отвечай на русском языке, опираясь "
        "ТОЛЬКО на приведённые фрагменты. Если информации недостаточно — скажи об этом. "
        "В конце ответа обязательно укажи источники в формате: "
        "Источники: [1] файл.pdf, [2] файл.docx"
    )
    user_prompt = (
        f"Вопрос пользователя:\n{question}\n\n"
        f"Фрагменты документов:\n{context_text}\n\n"
        "Дай точный и структурированный ответ со ссылками на источники."
    )

    headers = {
        "Authorization": f"Bearer {API_KEY}",
        "Content-Type": "application/json",
    }
    payload = {
        "model": DEEPSEEK_MODEL,
        "messages": [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
        "temperature": 0.2,
    }

    resp = requests.post(DEEPSEEK_URL, headers=headers, json=payload, timeout=120)
    resp.raise_for_status()
    data = resp.json()
    return data["choices"][0]["message"]["content"].strip()


def format_sources(contexts: List[Dict]) -> str:
    lines = ["\n--- Найденные фрагменты ---"]
    for i, c in enumerate(contexts, 1):
        preview = c["text"][:200] + ("..." if len(c["text"]) > 200 else "")
        lines.append(
            f"[{i}] {c['source']} | чанк #{c['chunk_id']} | score={c['score']:.3f}\n    {preview}"
        )
    return "\n".join(lines)


# =============================================================================
# Основной сценарий
# =============================================================================


def main():
    print("=" * 60)
    print("Система анализа документов (TF-IDF + DeepSeek)")
    print("Поддержка: PDF, DOCX, XLSX, TXT, сканы (OCR), ZIP, RAR, 7z")
    print("=" * 60)

    print("\nВыберите файлы для загрузки...")
    uploaded = files.upload()
    if not uploaded:
        print("Файлы не загружены.")
        return

    documents: List[Tuple[str, str]] = []
    print("\nИзвлечение текста...")
    for name, data in uploaded.items():
        path = WORK_DIR / _safe_name(name)
        path.write_bytes(data)
        print(f"\n• {name}")
        documents.extend(extract_from_file(path, data))

    if not documents:
        print("\nНе удалось извлечь текст ни из одного файла.")
        return

    print("\nПостроение TF-IDF индекса...")
    index = TfidfIndex()
    index.build(documents)

    print("\nГотово. Задавайте вопросы по документам.")
    print("Команды: пустая строка или 'exit' / 'выход' — завершить.\n")

    while True:
        try:
            question = input("Вопрос> ").strip()
        except EOFError:
            break

        if not question or question.lower() in {"exit", "quit", "выход", "q"}:
            print("Завершено.")
            break

        contexts = index.search(question, top_k=TOP_K)
        print(format_sources(contexts))

        print("\nЗапрос к DeepSeek...")
        try:
            answer = ask_deepseek(question, contexts)
        except Exception as e:
            print(f"Ошибка API DeepSeek: {e}")
            continue

        print("\n===== ОТВЕТ =====")
        print(answer)
        print("=================\n")


main()
