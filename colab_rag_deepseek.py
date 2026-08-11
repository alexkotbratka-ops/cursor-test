# =============================================================================
# RAG-система для Google Colab (DeepSeek + TF-IDF)
# Скопируйте весь код в одну ячейку Colab и запустите.
# =============================================================================

# --- 1. Установка системных пакетов и библиотек ---
import subprocess
import sys


def _run(cmd, check=True):
    print(f"$ {' '.join(cmd)}")
    subprocess.run(cmd, check=check)


print("Установка системных пакетов (OCR и архивы)...")
_run(
    [
        "apt-get",
        "update",
        "-qq",
    ],
    check=False,
)
_run(
    [
        "apt-get",
        "install",
        "-y",
        "-qq",
        "tesseract-ocr",
        "tesseract-ocr-rus",
        "tesseract-ocr-eng",
        "poppler-utils",
        "unrar",
        "p7zip-full",
    ],
    check=False,
)

print("Установка Python-библиотек...")
_run(
    [
        sys.executable,
        "-m",
        "pip",
        "install",
        "-q",
        "pymupdf",
        "python-docx",
        "openpyxl",
        "pytesseract",
        "pdf2image",
        "Pillow",
        "scikit-learn",
        "requests",
        "rarfile",
        "py7zr",
    ]
)

# --- Импорты ---
import getpass
import io
import os
import re
import shutil
import tempfile
import zipfile
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import fitz  # pymupdf
import openpyxl
import pytesseract
import rarfile
import requests
import py7zr
from docx import Document
from pdf2image import convert_from_path
from PIL import Image
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

try:
    from google.colab import files as colab_files
except ImportError:
    colab_files = None
    print(
        "Предупреждение: google.colab недоступен. "
        "Загрузка файлов через files.upload() работает только в Colab."
    )

# --- Константы ---
CHUNK_SIZE = 800
CHUNK_OVERLAP = 150
MIN_PDF_TEXT_CHARS = 80  # меньше — считаем PDF сканом
TOP_K = 5
DEEPSEEK_API_URL = "https://api.deepseek.com/chat/completions"
DEEPSEEK_MODEL = "deepseek-chat"
SESSION_KEY_ENV = "DEEPSEEK_API_KEY"

SUPPORTED_DOCS = {".pdf", ".docx", ".xlsx", ".txt"}
SUPPORTED_IMAGES = {".jpg", ".jpeg", ".png", ".tiff", ".tif", ".bmp"}
SUPPORTED_ARCHIVES = {".zip", ".rar", ".7z"}
ALL_SUPPORTED = SUPPORTED_DOCS | SUPPORTED_IMAGES | SUPPORTED_ARCHIVES


# =============================================================================
# 1. API-ключ (без хардкода)
# =============================================================================
def get_api_key() -> str:
    """Запрашивает DeepSeek API-ключ через getpass и сохраняет в переменную сессии."""
    key = os.environ.get(SESSION_KEY_ENV, "").strip()
    if key:
        print("API-ключ DeepSeek найден в переменных сессии.")
        return key

    print("Введите API-ключ DeepSeek (ввод скрыт):")
    key = getpass.getpass("DeepSeek API Key: ").strip()
    if not key:
        raise SystemExit(
            "Ошибка: API-ключ не введён. Система не запущена. "
            "Перезапустите ячейку и введите ключ."
        )

    os.environ[SESSION_KEY_ENV] = key
    print("API-ключ сохранён в переменную сессии.")
    return key


# =============================================================================
# 2. Извлечение текста из файлов
# =============================================================================
def extract_text_from_txt(path: Path) -> str:
    for encoding in ("utf-8", "cp1251", "latin-1"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_text(encoding="utf-8", errors="ignore")


def extract_text_from_docx(path: Path) -> str:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_text_from_xlsx(path: Path) -> str:
    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    parts = []
    for sheet in wb.worksheets:
        parts.append(f"[Лист: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            values = [str(c) for c in row if c is not None and str(c).strip()]
            if values:
                parts.append(" | ".join(values))
    wb.close()
    return "\n".join(parts)


def ocr_image(image: Image.Image, lang: str = "rus+eng") -> str:
    return pytesseract.image_to_string(image, lang=lang) or ""


def extract_text_from_image(path: Path) -> str:
    with Image.open(path) as img:
        # Конвертируем в RGB при необходимости (например, RGBA/P)
        if img.mode not in ("RGB", "L"):
            img = img.convert("RGB")
        return ocr_image(img)


def extract_text_from_pdf(path: Path) -> str:
    """Читает PDF; если мало текста (скан) — запускает OCR."""
    text_parts: List[str] = []
    with fitz.open(str(path)) as doc:
        for page in doc:
            text_parts.append(page.get_text("text") or "")
    text = "\n".join(text_parts).strip()

    if len(text) >= MIN_PDF_TEXT_CHARS:
        return text

    print(f"  PDF похож на скан ({len(text)} символов) — запускаю OCR: {path.name}")
    ocr_parts: List[str] = []
    try:
        images = convert_from_path(str(path), dpi=200)
        for i, image in enumerate(images, start=1):
            page_text = ocr_image(image)
            if page_text.strip():
                ocr_parts.append(page_text)
            print(f"    OCR страница {i}/{len(images)}")
    except Exception as e:
        print(f"  Ошибка OCR для {path.name}: {e}")
        return text
    return "\n".join(ocr_parts).strip() or text


def extract_text_from_file(path: Path) -> Optional[str]:
    ext = path.suffix.lower()
    try:
        if ext == ".txt":
            return extract_text_from_txt(path)
        if ext == ".docx":
            return extract_text_from_docx(path)
        if ext == ".xlsx":
            return extract_text_from_xlsx(path)
        if ext == ".pdf":
            return extract_text_from_pdf(path)
        if ext in SUPPORTED_IMAGES:
            print(f"  OCR изображения: {path.name}")
            return extract_text_from_image(path)
    except Exception as e:
        print(f"  Ошибка чтения {path.name}: {e}")
        return None
    return None


# =============================================================================
# Распаковка архивов (рекурсивно)
# =============================================================================
def _safe_extract_member(target_dir: Path, member_name: str) -> Optional[Path]:
    """Защита от path traversal при распаковке."""
    dest = (target_dir / member_name).resolve()
    if not str(dest).startswith(str(target_dir.resolve())):
        print(f"  Пропуск небезопасного пути в архиве: {member_name}")
        return None
    return dest


def extract_zip(archive: Path, dest: Path) -> None:
    with zipfile.ZipFile(archive, "r") as zf:
        for info in zf.infolist():
            if info.is_dir():
                continue
            out = _safe_extract_member(dest, info.filename)
            if out is None:
                continue
            out.parent.mkdir(parents=True, exist_ok=True)
            with zf.open(info) as src, open(out, "wb") as dst:
                shutil.copyfileobj(src, dst)


def extract_rar(archive: Path, dest: Path) -> None:
    with rarfile.RarFile(archive) as rf:
        for info in rf.infolist():
            if info.is_dir():
                continue
            out = _safe_extract_member(dest, info.filename)
            if out is None:
                continue
            out.parent.mkdir(parents=True, exist_ok=True)
            with rf.open(info) as src, open(out, "wb") as dst:
                shutil.copyfileobj(src, dst)


def extract_7z(archive: Path, dest: Path) -> None:
    with py7zr.SevenZipFile(archive, mode="r") as z:
        z.extractall(path=dest)


def unpack_archives_recursively(root: Path) -> None:
    """Рекурсивно распаковывает ZIP/RAR/7z внутри root."""
    changed = True
    while changed:
        changed = False
        archives = [
            p
            for p in root.rglob("*")
            if p.is_file() and p.suffix.lower() in SUPPORTED_ARCHIVES
        ]
        for archive in archives:
            extract_dir = archive.parent / f"_extracted_{archive.stem}"
            extract_dir.mkdir(parents=True, exist_ok=True)
            print(f"  Распаковка архива: {archive.name}")
            try:
                ext = archive.suffix.lower()
                if ext == ".zip":
                    extract_zip(archive, extract_dir)
                elif ext == ".rar":
                    extract_rar(archive, extract_dir)
                elif ext == ".7z":
                    extract_7z(archive, extract_dir)
                # Удаляем архив, чтобы не распаковывать повторно
                archive.unlink(missing_ok=True)
                changed = True
            except Exception as e:
                print(f"  Ошибка распаковки {archive.name}: {e}")


# =============================================================================
# Загрузка файлов и сбор документов
# =============================================================================
def upload_files(work_dir: Path) -> List[Path]:
    work_dir.mkdir(parents=True, exist_ok=True)
    if colab_files is None:
        raise SystemExit(
            "Ошибка: загрузка доступна только в Google Colab "
            "(google.colab.files.upload)."
        )

    print("Выберите файлы для загрузки (PDF, DOCX, XLSX, TXT, изображения, ZIP/RAR/7z)...")
    uploaded = colab_files.upload()
    if not uploaded:
        raise SystemExit("Ошибка: файлы не загружены. Система не запущена.")

    saved: List[Path] = []
    for name, content in uploaded.items():
        path = work_dir / Path(name).name
        path.write_bytes(content)
        saved.append(path)
        print(f"  Сохранён: {path.name} ({len(content)} байт)")
    return saved


def collect_source_files(root: Path) -> List[Path]:
    files = []
    for p in sorted(root.rglob("*")):
        if not p.is_file():
            continue
        if p.suffix.lower() in (SUPPORTED_DOCS | SUPPORTED_IMAGES):
            files.append(p)
    return files


def load_documents(work_dir: Path) -> Dict[str, str]:
    """
    Загружает файлы, распаковывает архивы, извлекает текст.
    Возвращает {имя_исходного_файла: текст}.
    """
    upload_files(work_dir)
    unpack_archives_recursively(work_dir)

    documents: Dict[str, str] = {}
    source_files = collect_source_files(work_dir)
    if not source_files:
        raise SystemExit("Ошибка: не найдено поддерживаемых документов после загрузки.")

    print(f"\nИзвлечение текста из {len(source_files)} файл(ов)...")
    for path in source_files:
        # Имя с относительным путём внутри work_dir, чтобы различать файлы из архивов
        rel_name = str(path.relative_to(work_dir))
        print(f"→ {rel_name}")
        text = extract_text_from_file(path)
        if text and text.strip():
            documents[rel_name] = text.strip()
            print(f"  Извлечено символов: {len(documents[rel_name])}")
        else:
            print(f"  Пустой результат, пропуск: {rel_name}")

    if not documents:
        raise SystemExit("Ошибка: не удалось извлечь текст ни из одного файла.")
    return documents


# =============================================================================
# 3. Чанки и TF-IDF индекс
# =============================================================================
def split_into_chunks(text: str, chunk_size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> List[str]:
    text = re.sub(r"\s+", " ", text).strip()
    if not text:
        return []
    if len(text) <= chunk_size:
        return [text]

    chunks = []
    start = 0
    step = max(1, chunk_size - overlap)
    while start < len(text):
        end = start + chunk_size
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= len(text):
            break
        start += step
    return chunks


class RAGIndex:
    def __init__(self):
        self.chunks: List[str] = []
        self.meta: List[Tuple[str, int]] = []  # (filename, chunk_number)
        self.vectorizer = TfidfVectorizer(
            analyzer="word",
            ngram_range=(1, 2),
            min_df=1,
            max_df=0.95,
        )
        self.matrix = None

    def build(self, documents: Dict[str, str]) -> None:
        self.chunks = []
        self.meta = []
        for filename, text in documents.items():
            parts = split_into_chunks(text)
            for i, chunk in enumerate(parts, start=1):
                self.chunks.append(chunk)
                self.meta.append((filename, i))
            print(f"  {filename}: {len(parts)} чанк(ов)")

        if not self.chunks:
            raise SystemExit("Ошибка: после нарезки нет чанков для индексации.")

        self.matrix = self.vectorizer.fit_transform(self.chunks)
        print(f"TF-IDF индекс построен: {len(self.chunks)} чанков, "
              f"словарь={len(self.vectorizer.vocabulary_)} токенов.")

    def search(self, query: str, top_k: int = TOP_K) -> List[dict]:
        if self.matrix is None or not self.chunks:
            return []
        q = self.vectorizer.transform([query])
        scores = cosine_similarity(q, self.matrix).ravel()
        if scores.max() <= 0:
            # fallback: вернуть первые top_k
            idxs = list(range(min(top_k, len(self.chunks))))
        else:
            idxs = scores.argsort()[::-1][:top_k]

        results = []
        for idx in idxs:
            fname, cnum = self.meta[idx]
            results.append(
                {
                    "text": self.chunks[idx],
                    "filename": fname,
                    "chunk_num": cnum,
                    "score": float(scores[idx]),
                }
            )
        return results


# =============================================================================
# 4. DeepSeek API и цикл вопросов
# =============================================================================
def ask_deepseek(api_key: str, question: str, contexts: List[dict]) -> str:
    context_blocks = []
    for i, c in enumerate(contexts, start=1):
        context_blocks.append(
            f"[{i}] Источник: {c['filename']}, чанк #{c['chunk_num']}\n{c['text']}"
        )
    context_text = "\n\n".join(context_blocks)

    system_prompt = (
        "Ты — помощник по анализу документов. Отвечай на вопрос пользователя "
        "только на основе предоставленного контекста. Если ответа нет в контексте, "
        "прямо скажи об этом. Отвечай на языке вопроса пользователя."
    )
    user_prompt = (
        f"Контекст из документов:\n{context_text}\n\n"
        f"Вопрос: {question}\n\n"
        "Дай полный и точный ответ. В конце кратко укажи, на какие источники опирался."
    )

    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json",
    }
    payload = {
        "model": DEEPSEEK_MODEL,
        "messages": [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
        "temperature": 0.2,
    }

    resp = requests.post(DEEPSEEK_API_URL, headers=headers, json=payload, timeout=120)
    if resp.status_code != 200:
        raise RuntimeError(
            f"DeepSeek API ошибка {resp.status_code}: {resp.text[:500]}"
        )
    data = resp.json()
    return data["choices"][0]["message"]["content"].strip()


def qa_loop(api_key: str, index: RAGIndex) -> None:
    print("\n" + "=" * 60)
    print("Система готова. Задавайте вопросы по документам.")
    print("Для выхода введите: exit или выход")
    print("=" * 60)

    while True:
        try:
            question = input("\nВопрос> ").strip()
        except EOFError:
            print("\nКонец ввода. Выход.")
            break

        if not question:
            continue
        if question.lower() in {"exit", "выход", "quit", "q"}:
            print("Выход из системы. До свидания!")
            break

        contexts = index.search(question, top_k=TOP_K)
        if not contexts:
            print("Релевантные фрагменты не найдены.")
            continue

        print("\nНайденные источники:")
        for c in contexts:
            print(
                f"  • {c['filename']} | чанк #{c['chunk_num']} "
                f"(score={c['score']:.4f})"
            )

        print("\nЗапрос к DeepSeek...")
        try:
            answer = ask_deepseek(api_key, question, contexts)
        except Exception as e:
            print(f"Ошибка при обращении к DeepSeek API: {e}")
            continue

        print("\n--- Ответ ---")
        print(answer)
        print("\n--- Источники ---")
        for c in contexts:
            print(f"  • файл: {c['filename']}, чанк: #{c['chunk_num']}")


# =============================================================================
# Главный запуск
# =============================================================================
def main():
    print("=" * 60)
    print("RAG-система: загрузка документов + DeepSeek")
    print("=" * 60)

    api_key = get_api_key()

    work_dir = Path(tempfile.mkdtemp(prefix="rag_docs_"))
    try:
        documents = load_documents(work_dir)

        print("\nНарезка на чанки (800 / overlap 150) и построение TF-IDF...")
        index = RAGIndex()
        index.build(documents)

        qa_loop(api_key, index)
    finally:
        # Очистка временных файлов
        shutil.rmtree(work_dir, ignore_errors=True)


if __name__ == "__main__":
    main()
